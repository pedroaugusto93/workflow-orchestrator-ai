from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from workflow_ai.domain.models import CaseRecord


ALIASES = {
    "PROCESSO": "process_id",
    "NOME_CURSO": "title",
    "OBJETO": "object_text",
    "CNPJ_CPF_FORNECEDOR": "supplier_document",
    "CNPJ_FORNECEDOR": "supplier_document",
    "NOME_FORNECEDOR": "supplier_name",
    "VALOR": "value",
    "PRAZO_EXECUCAO": "execution_term",
    "DATA_INICIO": "start_date",
    "DATA_EMPENHO": "end_date",
    "ANO_EMPENHO": "commitment_year",
    "NUM_EMPENHO": "commitment_number",
    "VALOR_EMPENHO": "commitment_value",
    "processoLink": "source_link",
    "file_path": "file_path",
    "resp_cpf": "responsible_document",
    "resp_email": "responsible_email",
    "autoridade_nome": "authority_name",
    "autoridade_cpf": "authority_document",
    "autoridade_email": "authority_email",
    "data_ato": "act_date",
    "ordenador": "ordering_officer",
}


def load_records(path: str | Path, sheet_name: str | None = None) -> list[CaseRecord]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        rows = ws.iter_rows(values_only=True)
        headers = [str(v or "").strip() for v in next(rows)]
        records: list[CaseRecord] = []
        for values in rows:
            raw = dict(zip(headers, values))
            known: dict[str, str] = {}
            extras = {}
            for key, value in raw.items():
                if not key:
                    continue
                target = ALIASES.get(key)
                if target:
                    known[target] = str(value or "").strip()
                else:
                    extras[key] = value
            if not known.get("process_id"):
                continue
            known["extras"] = extras
            records.append(CaseRecord(**known))
        return records
    finally:
        wb.close()


def group_by_process(records: Iterable[CaseRecord]) -> dict[str, list[CaseRecord]]:
    grouped: dict[str, list[CaseRecord]] = defaultdict(list)
    for record in records:
        grouped[record.process_id].append(record)
    return dict(grouped)
